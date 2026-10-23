import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="none"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def build_proposal_document(output_path: Path):
    doc = Document()

    # 1. Page Setup: A4 Portrait, Margins 18mm for concise, elegant 2-3 page fit
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    # 2. Header and Footer
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("ĐỀ XUẤT ĐỀ TÀI & ĐẶC TẢ CHỨC NĂNG SƠ BỘ — DỰ ÁN QUESY")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(128, 128, 128)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Quesy Assessment Platform • Chuẩn Thông tư 32/2018/TT-BGDĐT & Quyết định 764/QĐ-BGDĐT")
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(8.5)
    frun.font.italic = True
    frun.font.color.rgb = RGBColor(128, 128, 128)

    # Base Style Config
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(11)
    style_normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(3)

    # -------------------------------------------------------------
    # HEADER BANNER (Cơ quan chủ quản & Quốc hiệu)
    # -------------------------------------------------------------
    tbl_top = doc.add_table(rows=1, cols=2)
    tbl_top.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_top.autofit = False
    tbl_top.columns[0].width = Mm(85)
    tbl_top.columns[1].width = Mm(85)

    c0 = tbl_top.cell(0, 0)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p0.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    r.font.bold = True
    r.font.size = Pt(9.5)
    r = p0.add_run("TRƯỜNG ĐẠI HỌC / KHOA CÔNG NGHỆ THÔNG TIN\n")
    r.font.bold = True
    r.font.size = Pt(9.5)
    r = p0.add_run("BỘ MÔN CÔNG NGHỆ PHẦN MỀM")
    r.font.size = Pt(9)

    c1 = tbl_top.cell(0, 1)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r.font.bold = True
    r.font.size = Pt(9.5)
    r = p1.add_run("Độc lập – Tự do – Hạnh phúc\n")
    r.font.bold = True
    r.font.size = Pt(9.5)
    r = p1.add_run("-------------------o0o-------------------")
    r.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # DOCUMENT MAIN TITLE (Theo đúng hướng dẫn Project Proposal)
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(2)
    r_t1 = p_title.add_run("ĐỀ XUẤT ĐỀ TÀI & ĐẶC TẢ CHỨC NĂNG SƠ BỘ\n")
    r_t1.font.size = Pt(15)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(0, 51, 102) # Deep Navy Blue

    r_t2 = p_title.add_run("PROJECT PROPOSAL & INITIAL FUNCTIONAL REQUIREMENTS")
    r_t2.font.size = Pt(11)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(0, 113, 227) # Tech Blue

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("(Học phần: Công nghệ Phần mềm / Đồ án Môn học / Khóa luận tốt nghiệp)")
    r_sub.font.italic = True
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(100, 100, 100)

    # -------------------------------------------------------------
    # MỤC I. DANH SÁCH THÀNH VIÊN NHÓM (Chuẩn theo format người dùng yêu cầu)
    # -------------------------------------------------------------
    p_s1 = doc.add_paragraph()
    r = p_s1.add_run("I. DANH SÁCH THÀNH VIÊN:")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    p_s1.paragraph_format.space_before = Pt(6)
    p_s1.paragraph_format.space_after = Pt(3)

    members = [
        "1. [MSSV] - [Họ và tên] (Nhóm trưởng - Phụ trách Kiến trúc Hệ thống & Fullstack)",
        "2. [MSSV] - [Họ và tên] (Thành viên - Phụ trách Backend API & AI Microservices)",
        "3. [MSSV] - [Họ và tên] (Thành viên - Phụ trách Frontend SPA & UI/UX Design)",
        "4. [MSSV] - [Họ và tên] (Thành viên - Phụ trách Kiểm thử QA/Test & Chuẩn hóa Khảo thí)"
    ]

    for m in members:
        pm = doc.add_paragraph()
        pm.paragraph_format.space_after = Pt(2)
        pm.paragraph_format.left_indent = Mm(6)
        r = pm.add_run(m)
        r.font.size = Pt(11)
        if "[MSSV]" in m:
            r.font.italic = True

    # Giảng viên hướng dẫn
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    p_gv = doc.add_paragraph()
    p_gv.paragraph_format.left_indent = Mm(6)
    p_gv.paragraph_format.space_after = Pt(6)
    r_gv_label = p_gv.add_run("* Giảng viên hướng dẫn: ")
    r_gv_label.font.bold = True
    r_gv_label.font.size = Pt(10.5)
    r_gv_val = p_gv.add_run("[Học hàm/Học vị, Họ và tên Giảng viên hướng dẫn]")
    r_gv_val.font.italic = True
    r_gv_val.font.size = Pt(10.5)
    r_gv_val.font.color.rgb = RGBColor(90, 90, 90)

    # -------------------------------------------------------------
    # MỤC II. TÊN PHẦN MỀM (Chuẩn theo format outline của BA)
    # -------------------------------------------------------------
    p_s2 = doc.add_paragraph()
    r = p_s2.add_run("II. TÊN PHẦN MỀM:")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    p_s2.paragraph_format.space_before = Pt(6)
    p_s2.paragraph_format.space_after = Pt(3)

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.left_indent = Mm(6)
    p_p1.paragraph_format.space_after = Pt(2)
    r = p_p1.add_run("• Tên thương hiệu / Viết tắt: ")
    r.font.bold = True
    r2 = p_p1.add_run("QUESY (Question Ecosystem System)")
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0, 113, 227)

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.left_indent = Mm(6)
    p_p2.paragraph_format.space_after = Pt(2)
    r = p_p2.add_run("• Tên tiếng Việt: ")
    r.font.bold = True
    p_p2.add_run("Hệ thống Số hóa Đề thi bằng AI, Quản lý Ngân hàng Câu hỏi Chuẩn Quốc gia & Soạn thảo Đề thi LaTeX Tích hợp Trợ lý Trí tuệ Nhân tạo (AI Copilot Studio).")

    p_p3 = doc.add_paragraph()
    p_p3.paragraph_format.left_indent = Mm(6)
    p_p3.paragraph_format.space_after = Pt(2)
    r = p_p3.add_run("• Tên tiếng Anh: ")
    r.font.bold = True
    p_p3.add_run("Quesy — AI-Powered National Assessment Digitization, Item Banking & LaTeX Exam Composer Platform.")

    p_p4 = doc.add_paragraph()
    p_p4.paragraph_format.left_indent = Mm(6)
    p_p4.paragraph_format.space_after = Pt(3)
    p_p4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_p4.add_run("• Mô tả ngắn gọn: ")
    r.font.bold = True
    p_p4.add_run("Quesy là nền tảng phần mềm Web App chuyên sâu hỗ trợ giáo viên và các hội đồng khảo thí số hóa đề thi tự động từ tài liệu thô (PDF scan, Ảnh chụp, Word), giải mã công thức toán sang LaTeX, quản lý ngân hàng câu hỏi chuẩn hóa theo Thông tư 32/2018/TT-BGDĐT, thiết lập ma trận rút đề và xáo trộn mã đề an toàn theo định dạng đề thi mới từ năm 2025 (Quyết định 764/QĐ-BGDĐT), đồng thời cung cấp môi trường soạn đề LaTeX 6:4 có trợ lý AI hỗ trợ sửa đúng dòng và xuất bản đề thi chuẩn in ấn dạng vector PDF chất lượng cao.")

    p_p5 = doc.add_paragraph()
    p_p5.paragraph_format.left_indent = Mm(6)
    p_p5.paragraph_format.space_after = Pt(6)
    r = p_p5.add_run("• Nền tảng công nghệ: ")
    r.font.bold = True
    p_p5.add_run("Web Application (Single Page Application - SPA) kiến trúc hướng dịch vụ (FastAPI / Python 3.13, Node.js MathJax daemon, Docker Tectonic XeTeX Engine, HTML5/Tailwind CSS/PDF.js).")

    # -------------------------------------------------------------
    # MỤC III. CÁC YÊU CẦU CHỨC NĂNG CHÍNH CỦA PHẦN MỀM (THEO TƯ DUY BA)
    # -------------------------------------------------------------
    p_s3 = doc.add_paragraph()
    r = p_s3.add_run("III. CÁC YÊU CẦU CHỨC NĂNG CHÍNH CỦA PHẦN MỀM (MAIN FUNCTIONAL REQUIREMENTS):")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    p_s3.paragraph_format.space_before = Pt(8)
    p_s3.paragraph_format.space_after = Pt(4)

    # 1. Phân hệ Giáo viên / Người dùng chuyên môn
    p_act1 = doc.add_paragraph()
    r = p_act1.add_run("1. Phân hệ Người dùng chuyên môn / Giáo viên (Teacher / Expert User):")
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 204)
    p_act1.paragraph_format.space_before = Pt(4)
    p_act1.paragraph_format.space_after = Pt(2)

    f_act1 = [
        ("Quản lý tài khoản & Thiết lập môn học:", "Đăng nhập xác thực JWT/OAuth2, quản lý thông tin cá nhân, chọn môn học và phân hệ khối lớp giảng dạy (Toán, Lí, Hóa, Sinh... Lớp 10, 11, 12)."),
        ("Số hóa đề thi tự động bằng AI (AI Ingestion & OCR):", "Tải lên đề thi dạng file Word (.docx, .doc), PDF đa trang hoặc ảnh chụp (.png, .jpg); hệ thống tự động bóc tách cấu trúc câu hỏi, giải mã nhị phân MathType OLE (MTEF) sang LaTeX chuẩn, tự động nhận diện và cắt ảnh đồ thị, hình vẽ hình học gán định danh [IMAGE:url]."),
        ("Studio Hậu kiểm Đối chiếu Song song (HITL Review Studio):", "Giao diện Dual-pane cuộn trang đồng bộ (<100ms) đối chiếu giữa tài liệu gốc và thẻ câu hỏi số hóa; thanh đo độ tin cậy AI (Confidence Barometer); trình biên tập nhanh tại chỗ cho phép đổi đáp án 1-click, sửa công thức LaTeX có Live Preview, và sắp xếp lại trang trước khi nhập kho."),
        ("Bộ Dựng Đồ thị TikZ & Chuẩn hóa Toán học:", "Biên dịch mã TikZ sang đồ họa vector SVG sắc nét; thuật toán Auto-clipping tự động khống chế khung hình tiệm cận chống méo đồ thị; chuẩn hóa \\left/\\right kích thước động, số thập phân chuẩn Việt Nam {,}, phân số \\dfrac và vectơ \\overrightarrow."),
        ("Soạn đề LaTeX Trực tiếp & Trợ lý Trí tuệ Nhân tạo (LaTeX Workspace & AI Copilot):", "Bố cục Split-view 6:4 (60% Code Editor dark mode hỗ trợ gói ex_test.sty v3.3; 40% Live PDF Preview vector); Trợ lý AI popup Bottom Sheet hỗ trợ đề xuất thay đổi đúng dòng chỉ định (AI Inline Diff Smart Replacer: + thêm, - bớt) kèm nút Đồng ý / Từ chối; trình xem PDF.js đồng bộ 100% tỷ lệ thu phóng (Zoom Sync giữa nút bấm và Ctrl+cuộn chuột)."),
        ("Quản lý Ngân hàng Câu hỏi Cá nhân & Chống trùng lặp AI:", "Hỗ trợ trọn vẹn 5 hình thức thi QĐ 764 (Trắc nghiệm 4 lựa chọn, Đúng/Sai 4 khẳng định, Trả lời ngắn, Chùm bài đọc, Tự luận); phân loại theo Cây kiến thức TT 32; tự động đo khoảng cách Cosine Similarity qua vector embeddings để cảnh báo và chặn lưu câu hỏi trùng lặp (>= 90%)."),
        ("Thiết kế Ma trận Đề thi 2 Chiều & Rút đề tự động (Matrix & Solver):", "Lưới ma trận trực quan (Mạch nội dung x 3 Mức độ nhận thức: Nhận biết, Thông hiểu, Vận dụng); kiểm tra tồn kho câu hỏi thời gian thực; động cơ giải ràng buộc tự động rút đề gốc tối ưu, loại trừ câu hỏi cũ."),
        ("Động cơ Trộn Đề Hoán vị An toàn (Safe Shuffling Engine):", "Hoán vị tạo N mã đề thi (101, 102...) bảo đảm 4 nguyên tắc tuyệt đối: Khóa neo phương án phụ thuộc vị trí ở D (Option Anchor Rule), Bảo toàn bài đọc chung (Passage Integrity), Bảo toàn phân vùng phần thi (Section Integrity), và tự động xuất bảng đối chiếu đáp án chéo."),
        ("Xuất bản Tài liệu Khảo thí Chuẩn in ấn 1-Click:", "Xuất bản tài liệu đề thi và hướng dẫn giải dưới định dạng PDF vector độ phân giải cao chuẩn khổ giấy A4 Portrait, font Times New Roman / CMU Serif sắc nét, phiếu trả lời trắc nghiệm chuẩn máy chấm OMR và đóng gói trọn bộ tệp ZIP."),
        ("Chấm điểm Đúng/Sai Bậc thang & Quản lý Bảng điểm Lớp học:", "Tự động chấm điểm trắc nghiệm Đúng/Sai 4 ý theo thang điểm lũy tiến QĐ 764 (0.1đ – 0.25đ – 0.5đ – 1.0đ); quản lý danh sách học sinh theo SBD/Lớp; xuất/nhập bảng điểm Excel tương thích UTF-8 BOM chống lỗi font tiếng Việt.")
    ]

    for title, desc in f_act1:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Mm(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rt = p.add_run(f"{title} ")
        rt.font.bold = True
        rt.font.size = Pt(10.5)
        rd = p.add_run(desc)
        rd.font.size = Pt(10.5)

    # 2. Phân hệ Quản trị viên / Tổ trưởng chuyên môn
    p_act2 = doc.add_paragraph()
    r = p_act2.add_run("2. Phân hệ Quản trị viên / Tổ trưởng chuyên môn (Admin / Head of Department):")
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 204)
    p_act2.paragraph_format.space_before = Pt(6)
    p_act2.paragraph_format.space_after = Pt(2)

    f_act2 = [
        ("Quản lý Người dùng & Phân quyền (RBAC):", "Quản lý danh sách tài khoản giáo viên bộ môn, tổ trưởng, ban giám hiệu; phân quyền truy cập, chỉnh sửa và phê duyệt kho học liệu."),
        ("Quản lý Cây Kiến thức Chuẩn Quốc gia 6 Tầng (Taxonomy Management):", "Quản lý danh mục chuẩn hóa theo Thông tư 32/2018/TT-BGDĐT gồm 6 tầng: Môn học -> Khối lớp -> Mạch nội dung -> Chủ đề -> Bài học -> Yêu cầu cần đạt (YCCĐ) và Dạng câu hỏi chuyên đề. Tuyệt đối không phân chia theo các bộ sách giáo khoa thương mại."),
        ("Kiểm duyệt & Duyệt câu hỏi (Approval Workflow):", "Phê duyệt các câu hỏi đề xuất từ giáo viên vào ngân hàng câu hỏi dùng chung theo quy trình vòng đời: Bản nháp (DRAFT) -> Chờ duyệt (PENDING_REVIEW) -> Đã duyệt (APPROVED) -> Lưu trữ (ARCHIVED)."),
        ("Báo cáo & Thống kê Khảo thí (Dashboard & Psychometrics):", "Bảng điều khiển trực quan thống kê số lượng câu hỏi theo mức độ nhận thức, chủ đề; tự động tính toán chỉ số độ khó thực nghiệm (p) và độ phân biệt (D) sau các kỳ thi, tự động gắn cờ cảnh báo câu hỏi chất lượng kém (D < 0.20) để rà soát.")
    ]

    for title, desc in f_act2:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Mm(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rt = p.add_run(f"{title} ")
        rt.font.bold = True
        rt.font.size = Pt(10.5)
        rd = p.add_run(desc)
        rd.font.size = Pt(10.5)

    # 3. Các tính năng nâng cao & Điểm cộng công nghệ
    p_act3 = doc.add_paragraph()
    r = p_act3.add_run("3. Các tính năng nâng cao & Công nghệ nổi bật (Điểm cộng kỹ thuật của dự án):")
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 204)
    p_act3.paragraph_format.space_before = Pt(6)
    p_act3.paragraph_format.space_after = Pt(2)

    f_act3 = [
        ("Động cơ Biên dịch LaTeX Hybrid 3 tầng:", "Kết nối Remote Tectonic XeTeX Engine Docker chuyên dụng (host armbian:8090) hỗ trợ gói lệnh ex_test.sty và tự động bọc thư viện tcolorbox xparse, dự phòng Local Tectonic binary và Headless Chromium vector pipeline."),
        ("Động cơ Xuất bản PDF Vector Siêu tốc & Tự động Dàn trang:", "Xuất bản trực tiếp tệp PDF in ấn chuẩn quốc tế từ LaTeX và HTML5/ReportLab (< 100ms) với font Unicode tiếng Việt đầy đủ, loại bỏ phụ thuộc vào các phần mềm thương mại độc quyền."),
        ("AI Vector Embeddings Deduplication:", "Mô hình embedding đo khoảng cách Cosine Similarity giúp phát hiện trùng lặp câu hỏi thông minh, ngăn chặn câu hỏi lặp lại trong cùng đề thi."),
        ("Trình xem trước PDF Vector Canvas (PDF.js) Đồng bộ Thu phóng 100%:", "Khắc phục triệt để lỗi mất đồng bộ của trình đọc mặc định, đồng bộ 100% giữa nút bấm thanh nổi và phím tắt Ctrl + cuộn chuột, tự động căn giữa trang giấy."),
        ("Trải nghiệm Người dùng Chuẩn Apple HIG & Zero-Scroll:", "Giao diện SPA tối giản, phản hồi xúc giác tức thì (active:scale-[0.97]), bố cục khép kín 100% viewport không cuộn ngoài, hỗ trợ URL sạch deep-linking không bị lỗi 404 khi F5.")
    ]

    for title, desc in f_act3:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Mm(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rt = p.add_run(f"{title} ")
        rt.font.bold = True
        rt.font.size = Pt(10.5)
        rd = p.add_run(desc)
        rd.font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # SIGNATURE BLOCK
    # -------------------------------------------------------------
    tbl_sign = doc.add_table(rows=1, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sign.columns[0].width = Mm(85)
    tbl_sign.columns[1].width = Mm(85)

    cs0 = tbl_sign.cell(0, 0)
    ps0 = cs0.paragraphs[0]
    ps0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ps0.add_run("Ý KIẾN CỦA GIẢNG VIÊN HƯỚNG DẪN\n")
    r.font.bold = True
    r.font.size = Pt(10.5)
    r = ps0.add_run("(Phê duyệt ý tưởng & phạm vi đề tài)\n\n\n\n\n")
    r.font.italic = True
    r.font.size = Pt(9)
    r = ps0.add_run("[Họ và tên Giảng viên]")
    r.font.bold = True
    r.font.italic = True
    r.font.size = Pt(10)

    cs1 = tbl_sign.cell(0, 1)
    ps1 = cs1.paragraphs[0]
    ps1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ps1.add_run("Hà Nội, ngày ...... tháng ...... năm 2026\n")
    r.font.italic = True
    r.font.size = Pt(9.5)
    r = ps1.add_run("ĐẠI DIỆN NHÓM SINH VIÊN (NHÓM TRƯỞNG)\n")
    r.font.bold = True
    r.font.size = Pt(10.5)
    r = ps1.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n")
    r.font.italic = True
    r.font.size = Pt(9)
    r = ps1.add_run("[Họ và tên Nhóm trưởng]")
    r.font.bold = True
    r.font.italic = True
    r.font.size = Pt(10)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Successfully generated proposal document at: {output_path}")

if __name__ == "__main__":
    out_file = Path("De_Xuat_De_Tai_Quesy.docx").resolve()
    if len(sys.argv) > 1:
        out_file = Path(sys.argv[1]).resolve()
    build_proposal_document(out_file)
