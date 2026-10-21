"""DOCX Export Service for Quesy Exam Management System.
Implements 2-column layout, examination headers, and native OMML equations.
"""
import io
import re
from typing import Optional
from pydantic import BaseModel
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls

from backend.schemas import VariantOutput, VariantItem, ItemFormat


class ExamDocxConfig(BaseModel):
    school_name: str = "SỞ GIÁO DỤC & ĐÀO TẠO"
    exam_title: str = "KỲ THI KHẢO SÁT CHẤT LƯỢNG HỌC KỲ"
    subject: str = "MÔN: TOÁN HỌC"
    duration_minutes: int = 90
    num_columns: int = 2


def set_section_columns(section, num_columns: int = 2, space_twips: int = 720):
    """Set multi-column layout on a docx section."""
    sectPr = section._sectPr
    cols = sectPr.xpath(".//w:cols")
    if cols:
        cols[0].set(qn("w:num"), str(num_columns))
        cols[0].set(qn("w:space"), str(space_twips))
    else:
        cols_xml = f'<w:cols {nsdecls("w")} w:num="{num_columns}" w:space="{space_twips}"/>'
        sectPr.append(parse_xml(cols_xml))


def append_omml_equation(paragraph, equation_text: str):
    """Append a native OMML equation (<m:oMath>) to a paragraph."""
    # Clean latex math delimiters if present
    clean_eq = equation_text.strip()
    if clean_eq.startswith("${") and clean_eq.endswith("}$"):
        clean_eq = clean_eq[2:-2].strip()
    elif clean_eq.startswith("$$") and clean_eq.endswith("$$"):
        clean_eq = clean_eq[2:-2].strip()
    elif clean_eq.startswith("$") and clean_eq.endswith("$"):
        clean_eq = clean_eq[1:-1].strip()
    if clean_eq.startswith("{") and clean_eq.endswith("}"):
        clean_eq = clean_eq[1:-1].strip()

    from backend.services.docx_exporter import convert_latex_to_omml
    omml_elem = convert_latex_to_omml(clean_eq)
    if omml_elem is not None:
        paragraph._p.append(omml_elem)
        return

    omml_xml = (
        f'<m:oMath {nsdecls("m")}>'
        f'<m:r><m:t>{clean_eq}</m:t></m:r>'
        f'</m:oMath>'
    )
    paragraph._p.append(parse_xml(omml_xml))


def add_paragraph_with_latex(paragraph, text: str, bold_prefix: Optional[str] = None):
    """Parse text containing inline LaTeX formulas (${...}$, $$...$$, or $...$) and render them with OMML equations."""
    if bold_prefix:
        run = paragraph.add_run(bold_prefix)
        run.bold = True

    if not text:
        return

    # Split by inline LaTeX delimiters ${...}$, $${...}$$, $$...$$, $...$
    tokens = re.split(r"(\$\$\{.*?\}\$\$|\$\$.*?\$\$|\$\\{.*?\\}\$|\$\{.*?\}\$|\$.*?\$)", text, flags=re.DOTALL)
    for token in tokens:
        if not token:
            continue
        stripped = token.strip()
        is_math = (
            (stripped.startswith("${") and stripped.endswith("}$"))
            or (stripped.startswith("$$") and stripped.endswith("$$"))
            or (stripped.startswith("$") and stripped.endswith("$"))
        )
        if is_math and len(stripped) > 2:
            append_omml_equation(paragraph, stripped)
        else:
            paragraph.add_run(token)


def generate_exam_docx(
    variant: VariantOutput,
    config: Optional[ExamDocxConfig] = None,
    output_path: Optional[str] = None,
) -> bytes:
    """Generate a high-quality exam paper .docx file.
    
    Features:
    - Formal examination header (School, Subject, Exam Title, Variant code)
    - Two-column exam questions layout
    - Native OMML equation rendering
    """
    if config is None:
        config = ExamDocxConfig()

    doc = Document()

    # 1. Page Format: Standard A4 Portrait (210 x 297 mm)
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # 2. Header Table (School info on left, Exam title and code on right)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell_left = table.cell(0, 0)
    cell_right = table.cell(0, 1)

    p_left = cell_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_school = p_left.add_run(f"{config.school_name.upper()}\n")
    r_school.bold = True
    r_school.font.size = Pt(10)
    r_official = p_left.add_run("ĐỀ THI CHÍNH THỨC\n(Đề thi gồm các trang)")
    r_official.font.size = Pt(9.5)
    r_official.italic = True

    p_right = cell_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_right.add_run(f"{config.exam_title.upper()}\n")
    r_title.bold = True
    r_title.font.size = Pt(10.5)
    r_sub = p_right.add_run(f"{config.subject}\n")
    r_sub.bold = True
    r_sub.font.size = Pt(10)
    r_time = p_right.add_run(
        f"Thời gian làm bài: {config.duration_minutes} phút (không kể thời gian giao đề)\n"
    )
    r_time.font.size = Pt(9)
    r_code = p_right.add_run(f"Mã đề thi: {variant.variant_code}")
    r_code.bold = True
    r_code.font.size = Pt(10)

    # Student info paragraph
    p_student = doc.add_paragraph()
    p_student.paragraph_format.space_before = Pt(6)
    p_student.paragraph_format.space_after = Pt(8)
    r_stud = p_student.add_run(
        "Họ và tên thí sinh: ................................................................ "
        "Số báo danh: ................................"
    )
    r_stud.font.size = Pt(9.5)
    r_stud.italic = True

    # Section divider
    divider = doc.add_paragraph()
    divider.paragraph_format.space_after = Pt(8)
    r_div = divider.add_run("―" * 50)
    r_div.font.size = Pt(8)
    divider.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 3. Add section break with 2-column layout for exam questions
    q_section = doc.add_section()
    set_section_columns(q_section, num_columns=config.num_columns, space_twips=720)

    # 4. Render questions
    for q in variant.questions:
        p_stem = doc.add_paragraph()
        p_stem.paragraph_format.space_before = Pt(4)
        p_stem.paragraph_format.space_after = Pt(2)
        add_paragraph_with_latex(p_stem, q.stem, bold_prefix=f"Câu {q.position}: ")

        if q.options:
            for opt in q.options:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Inches(0.2)
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(1)
                add_paragraph_with_latex(p_opt, opt.content, bold_prefix=f"{opt.label}. ")

        elif q.sub_statements:
            for sub in q.sub_statements:
                p_sub = doc.add_paragraph()
                p_sub.paragraph_format.left_indent = Inches(0.2)
                p_sub.paragraph_format.space_before = Pt(1)
                p_sub.paragraph_format.space_after = Pt(1)
                add_paragraph_with_latex(p_sub, sub.content, bold_prefix=f"{sub.label}) ")

        elif q.item_format == ItemFormat.SHORT_ANSWER:
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Inches(0.2)
            p_ans.paragraph_format.space_before = Pt(1)
            p_ans.paragraph_format.space_after = Pt(1)
            r_ans = p_ans.add_run("Đáp án: ........................................................")
            r_ans.italic = True

    # End marker
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(12)
    r_end = p_end.add_run("---------- HẾT ----------")
    r_end.bold = True
    r_end.font.size = Pt(10)

    # Save to file if output_path provided
    if output_path:
        doc.save(output_path)

    # Return bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
