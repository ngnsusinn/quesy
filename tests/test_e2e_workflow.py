import pytest
import io
import zipfile
from fastapi.testclient import TestClient
from backend.main import app

@pytest.fixture(scope="module")
def client():
    return TestClient(app)


def test_complete_end_to_end_assessment_workflow(client):
    """
    Comprehensive End-to-End Integration Journey:
    1. Document Ingestion (Upload DOCX exam).
    2. AI Layout & Parsing (Extract questions, LaTeX, options).
    3. Commit Extracted Questions to Item Bank.
    4. Blueprint Specification (Create ExamMatrix for YCCDs).
    5. Constraint Solving (Assemble MasterExam).
    6. Safe Shuffling (Produce Variants with Option Anchor & Passage Integrity).
    7. Multi-variant Publishing (Download ZIP bundle with Word DOCX and Answer Sheets).
    8. Psychometric Assessment (Submit student response dataset, compute p-value & D-index).
    """

    # --------------------------------------------------------------------------
    # Step 1: Document Ingestion (Upload exam file)
    # --------------------------------------------------------------------------
    # Create sample docx binary in-memory
    import docx
    doc = docx.Document()
    doc.add_heading("ĐỀ KIỂM TRA ĐỊNH KỲ TOÁN 12 - THÔNG TƯ 32", 0)
    doc.add_paragraph("Câu 1: Cho hàm số $y = x^3 - 3x^2 + 2$. Tìm điểm cực tiểu của hàm số.")
    doc.add_paragraph("A. $x = 2$\nB. $x = 0$\nC. $x = -1$\nD. Cả A và B đều đúng")
    doc.add_paragraph("Câu 2: Cho hàm số $y = f(x)$. Các khẳng định sau đúng hay sai?")
    doc.add_paragraph("a) Hàm số đồng biến trên R.\nb) Hàm số có 2 cực trị.\nc) min y = -5.\nd) Đồ thị có tiệm cận đứng.")
    doc.add_paragraph("BẢNG ĐÁP ÁN: 1.A | 2.a-D, b-D, c-S, d-D")

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_bytes = docx_buffer.getvalue()

    upload_res = client.post(
        "/api/v1/digitize/upload",
        files={"file": ("De_Khao_Sat_Lop12.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert upload_res.status_code in [201, 202]
    upload_data = upload_res.json()
    job_id = upload_data["job_id"]
    assert job_id is not None
    assert upload_data["status"] in ["PENDING", "PROCESSING", "EXTRACTED", "COMPLETED"]

    # --------------------------------------------------------------------------
    # Step 2: Poll AI Digitize Job
    # --------------------------------------------------------------------------
    job_res = client.get(f"/api/v1/digitize/jobs/{job_id}")
    assert job_res.status_code == 200
    job_data = job_res.json()
    assert job_data["id"] == job_id
    assert job_data["status"] in ["EXTRACTED", "COMPLETED"]
    assert job_data["confidence_score"] >= 0.80

    # --------------------------------------------------------------------------
    # Step 3: Fetch Taxonomy YCCDs & Commit Questions to Item Bank
    # --------------------------------------------------------------------------
    tree_res = client.get("/api/v1/taxonomy/tree")
    assert tree_res.status_code == 200
    tree = tree_res.json()
    strand = tree["strands"][0]
    topic = strand["topics"][0]
    lesson = topic["lessons"][0]
    yccd = lesson["learning_outcomes"][0]
    archetype = yccd["archetypes"][0] if yccd.get("archetypes") else None

    # Commit question to bank
    commit_payload = {
        "questions": [
            {
                "stem": "Cho hàm số $y = x^3 - 3x^2 + 2$. Điểm cực đại của hàm số là:",
                "item_format": "SINGLE_CHOICE",
                "cognitive_level": "RECOGNITION",
                "learning_outcome_id": yccd["id"],
                "problem_archetype_id": archetype["id"] if archetype else None,
                "options": [
                    {"label": "A", "content": "$x = 0$", "is_correct": True, "is_fixed": False, "order_index": 1},
                    {"label": "B", "content": "$x = 2$", "is_correct": False, "is_fixed": False, "order_index": 2},
                    {"label": "C", "content": "$x = 1$", "is_correct": False, "is_fixed": False, "order_index": 3},
                    {"label": "D", "content": "Tất cả các đáp án trên đều sai", "is_correct": False, "is_fixed": True, "order_index": 4}
                ],
                "explanation": "Đạo hàm y' = 3x^2 - 6x = 0 <=> x=0 (cực đại), x=2 (cực tiểu)."
            },
            {
                "stem": "Cho hàm số bậc ba $y = f(x)$ có đồ thị như hình vẽ. Xét tính đúng sai của các mệnh đề:",
                "item_format": "TRUE_FALSE_GROUP",
                "cognitive_level": "APPLICATION",
                "learning_outcome_id": yccd["id"],
                "problem_archetype_id": archetype["id"] if archetype else None,
                "sub_statements": [
                    {"label": "a", "content": "Hàm số đồng biến trên $(2; +\\infty)$.", "is_true": True, "explanation": "Đồ thị đi lên."},
                    {"label": "b", "content": "Hàm số đạt cực tiểu tại $x = 2$.", "is_true": True, "explanation": "Điểm đáy cực tiểu."},
                    {"label": "c", "content": "Giá trị lớn nhất của hàm số là 4.", "is_true": False, "explanation": "Hàm số tiến ra vô cực."},
                    {"label": "d", "content": "Phương trình $f(x)=0$ có đúng 3 nghiệm phân biệt.", "is_true": True, "explanation": "Cắt trục hoành tại 3 điểm."}
                ],
                "explanation": "Khảo sát hàm số qua đồ thị chuẩn."
            }
        ]
    }
    commit_res = client.post(f"/api/v1/digitize/jobs/{job_id}/commit", json=commit_payload)
    assert commit_res.status_code == 200
    commit_data = commit_res.json()
    assert commit_data["committed_count"] >= 2

    # --------------------------------------------------------------------------
    # Step 4: Create ExamMatrix and Check Inventory Availability
    # --------------------------------------------------------------------------
    matrix_payload = {
        "title": "Ma trận Khảo thí Tích hợp E2E",
        "subject_id": tree["subject_id"],
        "grade_id": tree["grade_id"],
        "duration_minutes": 45,
        "total_points": 10.0,
        "items": [
            {
                "section": "Phần I: Trắc nghiệm 4 lựa chọn",
                "learning_outcome_id": yccd["id"],
                "cognitive_level": "RECOGNITION",
                "item_format": "SINGLE_CHOICE",
                "quantity": 1,
                "points_per_item": 5.0,
                "order_index": 1
            },
            {
                "section": "Phần II: Trắc nghiệm Đúng/Sai 4 ý",
                "learning_outcome_id": yccd["id"],
                "cognitive_level": "APPLICATION",
                "item_format": "TRUE_FALSE_GROUP",
                "quantity": 1,
                "points_per_item": 5.0,
                "order_index": 2
            }
        ]
    }
    mat_res = client.post("/api/v1/matrices", json=matrix_payload)
    assert mat_res.status_code == 201
    matrix = mat_res.json()
    matrix_id = matrix["id"]

    # Check bank inventory availability
    avail_res = client.post(f"/api/v1/matrices/{matrix_id}/check-availability")
    assert avail_res.status_code == 200
    avail_data = avail_res.json()
    assert avail_data["is_sufficient"] is True
    assert avail_data["total_deficit"] == 0

    # --------------------------------------------------------------------------
    # Step 5: Solve Matrix Constraints and Generate MasterExam
    # --------------------------------------------------------------------------
    gen_res = client.post("/api/v1/exams/generate", json={
        "matrix_id": matrix_id,
        "title": "Đề thi Gốc Master E2E"
    })
    assert gen_res.status_code == 201
    master_exam = gen_res.json()
    master_id = master_exam["id"]
    assert master_exam["total_questions"] == 2
    assert len(master_exam["items"]) == 2

    # --------------------------------------------------------------------------
    # Step 6: Safe Shuffling Engine (Produce 4 Variants)
    # --------------------------------------------------------------------------
    shuf_res = client.post(f"/api/v1/exams/{master_id}/shuffle", json={
        "variant_count": 4,
        "variant_codes": ["101", "102", "103", "104"],
        "shuffle_sections": True,
        "shuffle_options": True
    })
    assert shuf_res.status_code == 200
    shuf_data = shuf_res.json()
    assert shuf_data["generated_variants_count"] == 4
    assert shuf_data["variant_codes"] == ["101", "102", "103", "104"]
    assert len(shuf_data["answer_key_matrix"]) == 2

    # Query DB to verify RULE-02 (Option Anchor Rule) in all variants
    from backend.database import SessionLocal
    from backend.models import ExamVariant
    from backend.services.shuffling_engine import is_anchor_option
    import json
    db_session = SessionLocal()
    try:
        variants = db_session.query(ExamVariant).filter(ExamVariant.master_exam_id == master_id).all()
        assert len(variants) == 4
        for variant in variants:
            for v_item in variant.items:
                q = v_item.question
                has_fixed = any(is_anchor_option(opt) for opt in q.options) if (q and q.options) else False
                if has_fixed and v_item.option_mapping:
                    mapping = v_item.option_mapping if isinstance(v_item.option_mapping, dict) else json.loads(v_item.option_mapping)
                    assert mapping.get("D") == "D"  # Fixed option D must strictly stay at D
    finally:
        db_session.close()

    # --------------------------------------------------------------------------
    # Step 7: Export Word DOCX & Master Answer Key in ZIP Bundle
    # --------------------------------------------------------------------------
    zip_res = client.get(f"/api/v1/exams/{master_id}/export-zip")
    assert zip_res.status_code == 200
    assert zip_res.headers.get("content-type") == "application/zip"
    
    # Unpack and verify bundle content
    zip_archive = zipfile.ZipFile(io.BytesIO(zip_res.content))
    filenames = zip_archive.namelist()
    assert "De_101.docx" in filenames
    assert "De_102.docx" in filenames
    assert "De_103.docx" in filenames
    assert "De_104.docx" in filenames
    assert "BangDapAn_TongHop.docx" in filenames
    assert "LoiGiai_101.docx" in filenames

    # Verify DOCX files are non-empty and well-formed
    for fn in filenames:
        assert zip_archive.getinfo(fn).file_size > 1000

    # --------------------------------------------------------------------------
    # Step 8: Post-Exam Psychometrics (Item Difficulty p & Discrimination D)
    # --------------------------------------------------------------------------
    q1_id = master_exam["items"][0]["question_id"]
    q2_id = master_exam["items"][1]["question_id"]

    analytics_payload = {
        "student_scores": [
            {"student_id": "STU_01", "total_score": 10.0, "answers": {q1_id: True, q2_id: True}},
            {"student_id": "STU_02", "total_score": 10.0, "answers": {q1_id: True, q2_id: True}},
            {"student_id": "STU_03", "total_score": 5.0, "answers": {q1_id: True, q2_id: False}},
            {"student_id": "STU_04", "total_score": 5.0, "answers": {q1_id: True, q2_id: False}},
            {"student_id": "STU_05", "total_score": 5.0, "answers": {q1_id: False, q2_id: True}},
            {"student_id": "STU_06", "total_score": 0.0, "answers": {q1_id: False, q2_id: False}},
            {"student_id": "STU_07", "total_score": 0.0, "answers": {q1_id: False, q2_id: False}},
            {"student_id": "STU_08", "total_score": 0.0, "answers": {q1_id: False, q2_id: False}},
        ]
    }
    analytics_res = client.post(f"/api/v1/exams/{master_id}/analytics", json=analytics_payload)
    assert analytics_res.status_code == 200
    analytics_data = analytics_res.json()
    assert "items" in analytics_data
    assert len(analytics_data["items"]) == 2
    for item_metric in analytics_data["items"]:
        assert "difficulty_p" in item_metric
        assert "discrimination_d" in item_metric
        assert 0.0 <= item_metric["difficulty_p"] <= 1.0
