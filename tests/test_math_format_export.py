import io
import docx
from pathlib import Path
from backend.services.docx_exporter import (
    clean_math_expr,
    add_math_text_to_paragraph,
    generate_digitized_exam_docx
)

def test_clean_math_expr():
    assert clean_math_expr(r"${\overrightarrow{IJ} = \dfrac{1}{2}\left(\overrightarrow{CD} + \overrightarrow{AD} + \overrightarrow{DB}\right)}$") == r"\overrightarrow{IJ} = \dfrac{1}{2}\left(\overrightarrow{CD} + \overrightarrow{AD} + \overrightarrow{DB}\right)"
    assert clean_math_expr(r"${u_1 = 3}$") == "u_1 = 3"
    assert clean_math_expr(r"$u_6 = -96$") == "u_6 = -96"
    assert clean_math_expr(r"$$x^2 + 1$$") == "x^2 + 1"
    assert clean_math_expr(r"{\alpha + \beta}") == r"\alpha + \beta"

def test_add_math_text_omml_and_mathtype_and_latex():
    formula = r"${\overrightarrow{IJ} = \dfrac{1}{2}\left(\overrightarrow{CD} + \overrightarrow{AD} + \overrightarrow{DB}\right)}$"
    
    # 1. OMML (MathML)
    doc_omml = docx.Document()
    p1 = doc_omml.add_paragraph("Mệnh đề: ")
    add_math_text_to_paragraph(p1, f"Cho tứ diện ABCD, {formula}.", math_format="omml")
    xml_omml = doc_omml._part.blob.decode("utf-8", errors="ignore")
    assert "<m:oMath" in xml_omml

    # 2. MathType OLE
    doc_mt = docx.Document()
    p2 = doc_mt.add_paragraph("Mệnh đề: ")
    add_math_text_to_paragraph(p2, f"Cho tứ diện ABCD, {formula}.", math_format="mathtype")
    xml_mt = doc_mt._part.blob.decode("utf-8", errors="ignore")
    assert "<w:object" in xml_mt
    assert 'ProgID="Equation.DSMT4"' in xml_mt

    # 3. LaTeX wrapped in ${...}$
    doc_latex = docx.Document()
    p3 = doc_latex.add_paragraph("Mệnh đề: ")
    add_math_text_to_paragraph(p3, f"Cho tứ diện ABCD, {formula}.", math_format="latex")
    assert formula in p3.text

def test_generate_digitized_exam_formats(tmp_path: Path):
    questions = [
        {
            "id": 1,
            "item_format": "SINGLE_CHOICE",
            "stem": r"Cho tứ diện $ABCD$, gọi $I, J$ lần lượt là trung điểm của $AB$ và $CD$. Đẳng thức nào sau đây đúng?",
            "options": [
                {"label": "A", "content": r"${\overrightarrow{IJ} = \dfrac{1}{2}\left(\overrightarrow{CD} + \overrightarrow{AD} + \overrightarrow{DB}\right)}$"},
                {"label": "B", "content": r"${u_1 = 3}$"},
                {"label": "C", "content": r"${u_6 = -96}$"},
                {"label": "D", "content": r"${\overrightarrow{AB} + \overrightarrow{BC} = \overrightarrow{AC}}$"}
            ]
        }
    ]

    # Default format (must be MathType OLE)
    out_default = tmp_path / "exam_default.docx"
    generate_digitized_exam_docx(questions, out_default)
    doc_def = docx.Document(str(out_default))
    assert "<w:object" in doc_def._part.blob.decode("utf-8", errors="ignore")
    assert 'ProgID="Equation.DSMT4"' in doc_def._part.blob.decode("utf-8", errors="ignore")

    # OMML
    out_omml = tmp_path / "exam_omml.docx"
    generate_digitized_exam_docx(questions, out_omml, math_format="omml")
    doc1 = docx.Document(str(out_omml))
    assert "<m:oMath" in doc1._part.blob.decode("utf-8", errors="ignore")

    # MathType
    out_mt = tmp_path / "exam_mt.docx"
    generate_digitized_exam_docx(questions, out_mt, math_format="mathtype")
    doc2 = docx.Document(str(out_mt))
    assert "<w:object" in doc2._part.blob.decode("utf-8", errors="ignore")

    # LaTeX wrapped in ${...}$
    out_lat = tmp_path / "exam_latex.docx"
    generate_digitized_exam_docx(questions, out_lat, math_format="latex")
    doc3 = docx.Document(str(out_lat))
    text3 = " ".join(p.text for p in doc3.paragraphs)
    for t in doc3.tables:
        for r in t.rows:
            for c in r.cells:
                text3 += " " + c.text
    assert r"${\overrightarrow{IJ} = \dfrac{1}{2}\left(\overrightarrow{CD} + \overrightarrow{AD} + \overrightarrow{DB}\right)}$" in text3
    assert "${u_1 = 3}$" in text3
