import io
import os
import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.schemas import (
    Question,
    Option,
    CognitiveLevel,
    ItemFormat,
    VariantItem,
    VariantOutput,
)
from backend.services.docx_export import (
    generate_exam_docx,
    ExamDocxConfig,
)


def test_docx_export_2_columns_header_and_equations(tmp_path):
    """Test generating a .docx exam paper with 2 columns, header, and equations."""
    # Prepare sample exam questions with equations
    questions = [
        VariantItem(
            question_id="q-01",
            stem="Cho hàm số y = f(x) xác định trên R. Giá trị cực tiểu của hàm số $y = x^3 - 3x + 2$ là:",
            item_format=ItemFormat.SINGLE_CHOICE,
            position=1,
            options=[
                Option(label="A", content="$y_{CT} = 0$", is_correct=True),
                Option(label="B", content="$y_{CT} = 4$", is_correct=False),
                Option(label="C", content="$y_{CT} = 2$", is_correct=False),
                Option(label="D", content="$y_{CT} = -2$", is_correct=False),
            ],
        ),
        VariantItem(
            question_id="q-02",
            stem="Tìm nguyên hàm của hàm số $f(x) = \\int e^{2x} dx$.",
            item_format=ItemFormat.SINGLE_CHOICE,
            position=2,
            options=[
                Option(label="A", content="$\\frac{1}{2} e^{2x} + C$", is_correct=True),
                Option(label="B", content="$2 e^{2x} + C$", is_correct=False),
                Option(label="C", content="$e^{2x} + C$", is_correct=False),
                Option(label="D", content="Cả A và B đều đúng", is_correct=False, is_fixed=True),
            ],
        ),
    ]

    variant = VariantOutput(
        variant_code="101",
        questions=questions,
        answer_keys={"1": "A", "2": "A"},
    )

    config = ExamDocxConfig(
        school_name="SỞ GD&ĐT HÀ NỘI - TRƯỜNG THPT CHUYÊN",
        exam_title="ĐỀ KHẢO SÁT CHẤT LƯỢNG HỌC KỲ I - NĂM HỌC 2025 - 2026",
        subject="MÔN: TOÁN HỌC - LỚP 12",
        duration_minutes=90,
        num_columns=2,
    )

    output_path = tmp_path / "De_Thi_101.docx"
    docx_bytes = generate_exam_docx(variant, config, output_path=str(output_path))

    assert os.path.exists(output_path)
    assert len(docx_bytes) > 0

    # Load and inspect the generated Word document
    doc = Document(str(output_path))

    # 1. Verify Header contains school name, subject, exam title, and variant code
    full_text = []
    # Check tables (header block usually uses a table for school / exam layout)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)
    for p in doc.paragraphs:
        full_text.append(p.text)

    combined_text = " ".join(full_text)
    assert "SỞ GD&ĐT HÀ NỘI" in combined_text
    assert "TOÁN HỌC" in combined_text
    assert "101" in combined_text
    assert "Mã đề" in combined_text

    # 2. Verify 2-column layout in section properties
    # In OpenXML, column definition is w:cols in sectPr
    has_two_columns = False
    for section in doc.sections:
        cols_elements = section._sectPr.xpath(".//w:cols")
        if cols_elements:
            col_num = cols_elements[0].get(qn("w:num"))
            if col_num == "2":
                has_two_columns = True
                break

    assert has_two_columns, "Document does not have a 2-column section layout!"

    # 3. Verify presence of OMML native equations (<m:oMath>)
    omml_equations = doc._body._element.xpath(".//m:oMath")
    assert len(omml_equations) > 0, (
        f"Expected at least one native OMML equation (<m:oMath>), found {len(omml_equations)}."
    )

    # 4. Verify question stems and options are in the document
    assert "Câu 1" in combined_text
    assert "Câu 2" in combined_text
    assert "Cả A và B đều đúng" in combined_text
