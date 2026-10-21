from fastapi.testclient import TestClient
from backend.main import app

def test_export_digitized_docx_endpoint():
    client = TestClient(app)
    payload = {
        "title": "Đề Khảo Sát Toán 12",
        "questions": [
            {
                "id": 1,
                "item_format": "SINGLE_CHOICE",
                "stem": "Cho hàm số $y = \\dfrac{2x-1}{x+1}$. Tìm tiệm cận đứng.",
                "options": [
                    {"label": "A", "content": "$x = -1$"},
                    {"label": "B", "content": "$x = 2$"},
                    {"label": "C", "content": "$y = 2$"},
                    {"label": "D", "content": "$y = -1$"}
                ]
            },
            {
                "id": 2,
                "item_format": "SHORT_ANSWER",
                "stem": "Tính giá trị nhỏ nhất của $f(x) = x^2 - 4x + 5$ trên $[0; 3]$.",
                "short_answer_value": "1"
            }
        ]
    }
    response = client.post("/api/v1/digitize/export-docx", json=payload)
    assert response.status_code == 200
    assert response.headers.get("content-type") == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert len(response.content) > 1000
    assert response.content[:4] == b"PK\x03\x04"

    # Verify single-column layout and A4 Portrait paper format (210 x 297 mm)
    import io, docx
    from docx.enum.section import WD_ORIENT
    doc = docx.Document(io.BytesIO(response.content))
    assert len(doc.sections) == 1, "Exported exam must be 1 single column from top to bottom"
    s0 = doc.sections[0]
    assert s0.orientation == WD_ORIENT.PORTRAIT, "Must be Portrait orientation"
    assert abs(s0.page_width.mm - 210.0) < 1.0, f"Expected A4 width 210mm, got {s0.page_width.mm}mm"
    assert abs(s0.page_height.mm - 297.0) < 1.0, f"Expected A4 height 297mm, got {s0.page_height.mm}mm"
    sectPr = s0._sectPr
    cols = sectPr.xpath("./w:cols")
    num_cols = cols[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num") if cols else "1"
    assert num_cols in (None, "1"), f"Expected 1 column, got {num_cols}"

    # Verify formulas are embedded as genuine MathType OLE objects (<w:object>)
    xml_str = doc._part.blob.decode("utf-8", errors="ignore")
    assert "<w:object" in xml_str, "Formulas must be exported as genuine MathType OLE objects!"
    assert 'ProgID="Equation.DSMT4"' in xml_str, "Must have MathType Equation.DSMT4 ProgID!"

    # Also test explicit math_format="latex" to verify ${...}$ wrapping
    payload_latex = dict(payload, math_format="latex")
    resp_latex = client.post("/api/v1/digitize/export-docx", json=payload_latex)
    assert resp_latex.status_code == 200
    doc_latex = docx.Document(io.BytesIO(resp_latex.content))
    full_text_latex = " ".join(p.text for p in doc_latex.paragraphs)
    for table in doc_latex.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text_latex += " " + cell.text

    assert "${y = \\dfrac{2x-1}{x+1}}$" in full_text_latex
    assert "${x = -1}$" in full_text_latex
    assert "${f(x) = x^2 - 4x + 5}$" in full_text_latex
